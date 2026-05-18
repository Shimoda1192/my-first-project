import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "QDC_FULL基本機能確認"

HOURS_PER_DAY = 8

# Color palette
COLOR_HEADER_DARK  = "1F3864"  # dark navy
COLOR_HEADER_MID   = "2E75B6"  # blue
COLOR_SECTION_BG   = "D6E4F0"  # light blue
COLOR_SUBSEC_BG    = "EBF3FB"  # lighter blue
COLOR_ALT_ROW      = "F2F9FF"  # very light blue
COLOR_TOTAL_BG     = "FFF2CC"  # yellow
COLOR_WHITE        = "FFFFFF"
COLOR_MANUAL_H     = "C6EFCE"  # light green
COLOR_MANUAL_D     = "E2EFDA"  # pale green
COLOR_AUTO_H       = "FFEB9C"  # light amber
COLOR_AUTO_D       = "FFF2CC"  # pale amber

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def thick_bottom():
    thin = Side(style="thin", color="AAAAAA")
    thick = Side(style="medium", color="333333")
    return Border(left=thin, right=thin, top=thin, bottom=thick)

def font(bold=False, color="000000", size=10):
    return Font(bold=bold, color=color, size=size, name="Meiryo")

def align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

# ------------------------------------------------------------------
# Column layout
# A: 大分類, B: 中分類, C: 詳細項目
# D: 手動操作(h), E: 手動操作(日), F: 自動操作(h), G: 自動操作(日)
# ------------------------------------------------------------------

ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 22
ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 12
ws.column_dimensions["F"].width = 14
ws.column_dimensions["G"].width = 12

# ---- Row 1: Title ------------------------------------------------
ws.merge_cells("A1:G1")
c = ws["A1"]
c.value = "QDC_FULL 基本機能確認（QFC, QDC）"
c.font = Font(bold=True, color=COLOR_WHITE, size=13, name="Meiryo")
c.fill = fill(COLOR_HEADER_DARK)
c.alignment = align("center", "center")
c.border = thin_border()
ws.row_dimensions[1].height = 26

# ---- Row 2: Column headers ---------------------------------------
headers = [
    ("大分類", "A2"), ("中分類", "B2"), ("詳細項目", "C2"),
    ("手動操作\n(h)", "D2"), ("手動操作\n(日)", "E2"),
    ("自動操作\n(h)", "F2"), ("自動操作\n(日)", "G2"),
]
for label, addr in headers:
    c = ws[addr]
    c.value = label
    c.font = Font(bold=True, color=COLOR_WHITE, size=10, name="Meiryo")
    c.fill = fill(COLOR_HEADER_MID)
    c.alignment = align("center", "center", wrap=True)
    c.border = thin_border()
ws.row_dimensions[2].height = 32

# ------------------------------------------------------------------
# Data rows
# Format: (大分類, 中分類, 詳細項目, 手動h, 自動h)
# ------------------------------------------------------------------
data = [
    # 実験時間 / 実験準備時間
    ("実験時間", "実験準備時間(h)",
     "データ整理\n（UPロード時間、データ確認）", 1.2, 1.2),
    ("",         "",
     "スタンバイ\n（C/D載せ・降ろし）",           3.0, 3.0),
    ("",         "",
     "計測準備\n（設置、設定）",                   9.5, 9.5),
    ("",         "",
     "DTC取得",                                     1.2, 1.2),
    # 実験時間 / 実験評価時間
    ("",         "実験評価時間(h)",
     "CD実験・作業",                                8.0, 8.0),
    ("",         "",
     "定置実験・作業",                              6.0, 6.0),
    # 解析時間
    ("解析時間", "波形作成、クライテリア確認時間(h)",
     "",                                            16.0, 16.0),
]

# rows start at 3
ROW_START = 3

# Track merge ranges for 大分類 and 中分類
section_ranges = {}   # key: (col, value) -> [first_row, last_row]

for i, (cat1, cat2, detail, manual_h, auto_h) in enumerate(data):
    row = ROW_START + i
    ws.row_dimensions[row].height = 36

    bg = COLOR_ALT_ROW if i % 2 == 0 else COLOR_WHITE
    if cat1 == "解析時間":
        bg = COLOR_SUBSEC_BG

    # A: 大分類
    ca = ws.cell(row=row, column=1, value=cat1 if cat1 else None)
    ca.font = font(bold=bool(cat1), size=10)
    ca.fill = fill(COLOR_SECTION_BG if cat1 else bg)
    ca.alignment = align("center", "center", wrap=True)
    ca.border = thin_border()

    # B: 中分類
    cb = ws.cell(row=row, column=2, value=cat2 if cat2 else None)
    cb.font = font(bold=bool(cat2), size=10)
    cb.fill = fill(COLOR_SUBSEC_BG if cat2 else bg)
    cb.alignment = align("center", "center", wrap=True)
    cb.border = thin_border()

    # C: 詳細項目
    cc = ws.cell(row=row, column=3, value=detail if detail else None)
    cc.font = font(size=10)
    cc.fill = fill(bg)
    cc.alignment = align("left", "center", wrap=True)
    cc.border = thin_border()

    # D: 手動h
    cd = ws.cell(row=row, column=4, value=manual_h)
    cd.font = font(size=10)
    cd.fill = fill(COLOR_MANUAL_H)
    cd.alignment = align("center", "center")
    cd.number_format = "0.0"
    cd.border = thin_border()

    # E: 手動日
    ce = ws.cell(row=row, column=5, value=round(manual_h / HOURS_PER_DAY, 2))
    ce.font = font(size=10)
    ce.fill = fill(COLOR_MANUAL_D)
    ce.alignment = align("center", "center")
    ce.number_format = "0.00"
    ce.border = thin_border()

    # F: 自動h
    cf = ws.cell(row=row, column=6, value=auto_h)
    cf.font = font(size=10)
    cf.fill = fill(COLOR_AUTO_H)
    cf.alignment = align("center", "center")
    cf.number_format = "0.0"
    cf.border = thin_border()

    # G: 自動日
    cg = ws.cell(row=row, column=7, value=round(auto_h / HOURS_PER_DAY, 2))
    cg.font = font(size=10)
    cg.fill = fill(COLOR_AUTO_D)
    cg.alignment = align("center", "center")
    cg.number_format = "0.00"
    cg.border = thin_border()

# Merge 大分類: rows 3-8 (実験時間), row 9 (解析時間)
ws.merge_cells(f"A{ROW_START}:A{ROW_START+5}")
ws.merge_cells(f"A{ROW_START+6}:A{ROW_START+6}")
# Merge 中分類: rows 3-6 (実験準備時間), rows 7-8 (実験評価時間)
ws.merge_cells(f"B{ROW_START}:B{ROW_START+3}")
ws.merge_cells(f"B{ROW_START+4}:B{ROW_START+5}")
# 解析時間 中分類 spans C as well (no detail column)
ws.merge_cells(f"B{ROW_START+6}:C{ROW_START+6}")

# ---- TOTAL row ---------------------------------------------------
TOTAL_ROW = ROW_START + len(data)
ws.row_dimensions[TOTAL_ROW].height = 28

total_manual_h = sum(r[3] for r in data)
total_auto_h   = sum(r[4] for r in data)

ws.merge_cells(f"A{TOTAL_ROW}:C{TOTAL_ROW}")
ct = ws[f"A{TOTAL_ROW}"]
ct.value = "TOTAL"
ct.font = Font(bold=True, color="000000", size=11, name="Meiryo")
ct.fill = fill(COLOR_TOTAL_BG)
ct.alignment = align("center", "center")
ct.border = thick_bottom()

for col, val, fmt in [
    (4, total_manual_h,                      "0.0"),
    (5, round(total_manual_h/HOURS_PER_DAY, 2), "0.00"),
    (6, total_auto_h,                        "0.0"),
    (7, round(total_auto_h/HOURS_PER_DAY, 2),   "0.00"),
]:
    c = ws.cell(row=TOTAL_ROW, column=col, value=val)
    c.font = Font(bold=True, size=11, name="Meiryo")
    c.fill = fill(COLOR_TOTAL_BG)
    c.alignment = align("center", "center")
    c.number_format = fmt
    c.border = thick_bottom()

# ---- Legend row --------------------------------------------------
LEG_ROW = TOTAL_ROW + 2
ws.merge_cells(f"A{LEG_ROW}:B{LEG_ROW}")
ws[f"A{LEG_ROW}"].value = "凡例：1日 = 8時間換算"
ws[f"A{LEG_ROW}"].font = Font(italic=True, size=9, color="666666", name="Meiryo")
ws[f"A{LEG_ROW}"].alignment = align("left", "center")

ws.sheet_view.showGridLines = False

wb.save("/home/user/my-first-project/QDC_FULL_schedule.xlsx")
print("Excel saved.")
